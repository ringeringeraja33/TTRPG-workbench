"""Local-only document readers. No macros, HTML scripts or embedded objects execute.
Optional OCR is an extraction aid; its confidence is not a rule verification score.
"""
import json,re,zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from html.parser import HTMLParser
ocr=None

def recognize(im):
 global ocr
 if ocr is None:
  from rapidocr_onnxruntime import RapidOCR
  ocr=RapidOCR(det_use_cuda=False,cls_use_cuda=False,rec_use_cuda=False,intra_op_num_threads=4,inter_op_num_threads=2)
 import numpy as np
 rows,_=ocr(np.array(im.convert("RGB")))
 return {"text":"\n".join(r[1] for r in rows or []),"ocr_min_confidence":min((float(r[2]) for r in rows or []),default=0),"ocr_mean_confidence":sum(float(r[2]) for r in rows or [])/max(1,len(rows or [])),"method":"OCR; visual review pending"}

def decode(data):
 for enc in ("utf-8-sig","utf-16","gb18030","cp1252"):
  try:
   if enc=="utf-16" and not data.startswith((b"\xff\xfe",b"\xfe\xff")): continue
   return data.decode(enc),enc
  except UnicodeError: pass
 raise ValueError("No strict decoder")

def extract(p, *, chm_root=None, ocr_enabled=False):
 ext=p.suffix.lower()
 if ext==".pdf":
  from pypdf import PdfReader
  reader=PdfReader(p);render=None
  for i,page in enumerate(reader.pages):
   try: t=page.extract_text() or ""
   except Exception: t=""
   if len(re.sub(r"\s","",t))<25:
    if not ocr_enabled:
     yield {"anchor":f"PDF:{i+1}","text":t,"method":"PDF text layer; sparse page needs visual review or OCR"}
     continue
    import pypdfium2 as pdfium
    if render is None: render=pdfium.PdfDocument(p)
    u=recognize(render[i].render(scale=1.7).to_pil())
   else: u={"text":t,"method":"PDF text layer"}
   yield dict(anchor=f"PDF:{i+1}",**u)
  if render: render.close()
 elif ext==".docx":
  with zipfile.ZipFile(p) as z:
   for name in sorted(n for n in z.namelist() if re.fullmatch(r"word/(document|footnotes|endnotes|header\d+|footer\d+)\.xml",n)):
    root=ET.fromstring(z.read(name))
    for i,e in enumerate((e for e in root.iter() if e.tag.endswith("}p")),1):
     t="".join(x.text or "" for x in e.iter() if x.tag.endswith("}t"))
     if t.strip(): yield {"anchor":f"{name}:paragraph:{i}","text":t,"method":"OOXML"}
 elif ext==".xlsx":
  import openpyxl
  b=openpyxl.load_workbook(p,read_only=True,data_only=False);c=openpyxl.load_workbook(p,read_only=False,data_only=True)
  try:
   for sheet in b:
    for row in sheet:
     for cell in row:
      if cell.value is not None:
       yield {"anchor":f"sheet:{sheet.title}:cell:{cell.coordinate}","text":str(cell.value),"cached":c[sheet.title][cell.coordinate].value,"formula":cell.data_type=="f","method":"OOXML; no recalculation"}
  finally: b.close();c.close()
 elif ext in (".png",".jpg",".jpeg"):
  from PIL import Image
  with Image.open(p) as im:
   if ocr_enabled: yield dict(anchor="image:1",width=im.width,height=im.height,**recognize(im))
   else: yield {"anchor":"image:1","width":im.width,"height":im.height,"text":"","method":"image metadata; visual review pending"}
 elif ext==".txt":
  t,enc=decode(p.read_bytes())
  yield {"anchor":"text:1","text":t,"method":"strict "+enc}
 elif ext==".chm":
  if chm_root is None: raise ValueError("CHM needs a separately extracted --chm-root directory")
  d=Path(chm_root)/p.stem
  files=sorted(f for f in d.rglob("*") if f.is_file() and f.suffix.lower() in (".html",".htm",".hhc",".hhk",".txt"))
  if not files: raise ValueError("CHM extraction produced no documents")
  class Text(HTMLParser):
   def __init__(self): super().__init__();self.parts=[];self.skip=0
   def handle_starttag(self,tag,attrs):
    if tag in ("script","style"): self.skip+=1
   def handle_endtag(self,tag):
    if tag in ("script","style"): self.skip=max(0,self.skip-1)
   def handle_data(self,data):
    if not self.skip and data.strip(): self.parts.append(data.strip())
  for f in files:
   t,enc=decode(f.read_bytes());parser=Text();parser.feed(t)
   yield {"anchor":f"CHM:{f.relative_to(d).as_posix()}","text":"\n".join(parser.parts),"method":"HTML "+enc}
 elif ext==".doc":
  import olefile,struct
  with olefile.OleFileIO(p) as ole:
   w=ole.openstream("WordDocument").read()
   flags=struct.unpack_from("<H",w,10)[0];table=ole.openstream("1Table" if flags&0x200 else "0Table").read()
   csw=struct.unpack_from("<H",w,32)[0];off=34+csw*2;cslw=struct.unpack_from("<H",w,off)[0];off+=2+cslw*4
   count=struct.unpack_from("<H",w,off)[0];base=off+2
   if count<=33: raise ValueError("Unsupported legacy Word FIB")
   fc,lcb=struct.unpack_from("<II",w,base+33*8);clx=table[fc:fc+lcb];i=0
   while i<len(clx) and clx[i]==1: i+=3+struct.unpack_from("<H",clx,i+1)[0]
   if i>=len(clx) or clx[i]!=2: raise ValueError("Missing Word piece table")
   nbytes=struct.unpack_from("<I",clx,i+1)[0];plc=clx[i+5:i+5+nbytes];n=(nbytes-4)//12
   texts=[]
   for j in range(n):
    a,b=struct.unpack_from("<II",plc,j*4);pos=struct.unpack_from("<I",plc,4*(n+1)+j*8+2)[0];compressed=bool(pos&0x40000000);pos&=0x3fffffff
    if compressed: pos//=2
    data=w[pos:pos+(b-a)*(1 if compressed else 2)];texts.append(data.decode("cp1252" if compressed else "utf-16le"))
   yield {"anchor":"Word:body","text":"".join(texts).replace("\r","\n"),"method":"OLE Word piece table; layout not reviewed"}
 else: raise ValueError("Unsupported format")
